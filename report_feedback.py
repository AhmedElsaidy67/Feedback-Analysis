import os
import io
import pandas as pd
import streamlit as st
import plotly.express as px
import matplotlib.pyplot as plt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from shared_utils import save_master_roster

# --- Smart Pre-Classification Engine ---
def smart_pre_classify_row(row, hist_map=None):
    """
    Infers Handling Result ('Reasonable', 'Unreasonable', 'SNULL') and Confidence Level
    from system columns so the user doesn't have to open every Electronic Flow Link manually!
    """
    order_no = str(row.get("Order No.", "")).strip()
    if hist_map and order_no in hist_map:
        h = hist_map[order_no]
        return h["result"], "100% (From Previous Sheet)", h.get("comment", ""), h.get("severity", ""), h.get("escalation", "")

    sol1 = str(row.get("Responsible person replies with a handling solution", "SNULL")).strip()
    sol2 = str(row.get("Responsible person replies with a handling solution.1", sol1)).strip()
    resolved = str(row.get("Is it resolved?", "SNULL")).strip()
    loop_st = str(row.get("Has the responsible person promptly closed the loop?", "")).strip()
    upg_res = str(row.get("Upgrade of requirements, human handling of results", "SNULL")).strip()
    sys_res = str(row.get("The responsible person replies with the handling result.", "SNULL")).strip()
    reply_txt = str(row.get("Responsible Person's Reply and Comments", "SNULL")).strip()
    issue_type = str(row.get("Issue Type", "")).strip()

    # Rule 1: If already explicitly labeled Reasonable or Unreasonable in the file
    if sys_res in ["Reasonable", "Unreasonable"]:
        sev = "High" if (sys_res == "Unreasonable" or issue_type == "The content is incorrect.") else ("Medium" if issue_type == "Situation encountered by uncovered users" else "Low")
        esc = "Yes" if sys_res == "Unreasonable" else "No"
        return sys_res, "100% (Explicit in File)", "", sev, esc

    # Rule 2: If Handling Solution is Corrected Knowledge or Optimized knowledge -> 100% Reasonable!
    if sol1 in ["Corrected Knowledge", "Optimized knowledge"] or sol2 in ["Corrected Knowledge", "Optimized knowledge"] or upg_res == "The complaint is reasonable":
        return "Reasonable", "High (Auto: Knowledge Updated)", "", "Low", "No"

    # Rule 3: If Upgrade result says unreasonable OR Is it resolved? == 'The issue remains unresolved.' -> Unreasonable!
    if upg_res == "The complaint is unreasonable." or resolved == "The issue remains unresolved.":
        return "Unreasonable", "High (Auto: Unresolved / Rejected)", "", "High", "Yes"

    # Rule 4: Check handler reply keywords if solution is SNULL
    reply_lower = reply_txt.lower()
    if reply_txt != "SNULL" and reply_txt != "" and reply_txt != "nan":
        reject_phrases = ["sorry", "no changes", "not available", "will not", "already modified", "originally in english", "ineffective", "not a knowledge"]
        accept_phrases = ["done", "changed to", "updated", "added", "modified the knowledge", "received feedback and changed"]
        if any(p in reply_lower for p in reject_phrases):
            return "Unreasonable", "Medium-High (Auto: Rejection Reply)", "", "High", "Yes"
        if any(p in reply_lower for p in accept_phrases):
            return "Reasonable", "Medium-High (Auto: Positive Reply)", "", "Low", "No"

    # Rule 5: If Loop is still 'The electronic current has not been fully completed.' and no reply -> SNULL (Still Pending)
    if loop_st == "The electronic current has not been fully completed." and (reply_txt == "SNULL" or not reply_txt):
        sev = "High" if issue_type == "The content is incorrect." else ("Medium" if issue_type == "Situation encountered by uncovered users" else "Low")
        return "SNULL", "High (Auto: Loop Incomplete & No Reply)", "", sev, "Yes"

    # Rule 6: Remaining ambiguous cases (e.g. 'Not through the owner phase' with SNULL reply) -> Needs quick link check
    return "Needs Link Check", "Low (Check E-Flow Link)", "", "Medium", "No"


def enrich_and_classify_dataframe(df, agents_df, hist_map=None):
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    if "Consignee" not in df.columns:
        for col in df.columns:
            if col.lower() in ["consignee", "consignee id", "staff id", "agent id", "user id", "account"]:
                df.rename(columns={col: "Consignee"}, inplace=True)
                break
    if "Consignee" not in df.columns:
        st.error(f"❌ Column 'Consignee' was not found in the uploaded sheet! Columns found: {list(df.columns)}")
        st.stop()

    df["Consignee"] = df["Consignee"].astype(str).str.strip()

    roster_map_name = dict(zip(agents_df["Consignee"].str.upper(), agents_df["Submitter Name"]))
    roster_map_lob = dict(zip(agents_df["Consignee"].str.upper(), agents_df["Submitter LOB"]))

    mapped_names = df["Consignee"].str.upper().map(roster_map_name)
    mapped_lobs = df["Consignee"].str.upper().map(roster_map_lob)

    if "Submitter Name" in df.columns:
        df["Submitter Name"] = mapped_names.fillna(df["Submitter Name"].astype(str).str.strip())
    else:
        c_idx = df.columns.get_loc("Consignee")
        df.insert(c_idx + 1, "Submitter Name", mapped_names.fillna("Unknown Agent"))

    if "Submitter LOB" in df.columns:
        df["Submitter LOB"] = mapped_lobs.fillna(df["Submitter LOB"].astype(str).str.strip())
    else:
        n_idx = df.columns.get_loc("Submitter Name")
        df.insert(n_idx + 1, "Submitter LOB", mapped_lobs.fillna("Unassigned LOB"))

    df["Submitter Name"] = df["Submitter Name"].astype(str).str.strip()
    df["Submitter LOB"] = df["Submitter LOB"].astype(str).str.strip()

    # Apply smart pre-classification
    results, confs, comms, sevs, escs = [], [], [], [], []
    for _, row in df.iterrows():
        r_val, c_val, cm_val, sv_val, es_val = smart_pre_classify_row(row, hist_map)
        results.append(r_val)
        confs.append(c_val)
        comms.append(cm_val)
        sevs.append(sv_val)
        escs.append(es_val)

    res_col = "The responsible person replies with the handling result."
    df[res_col] = results
    df["Auto_Detection_Confidence"] = confs
    df["Comment"] = df["Comment"] if "Comment" in df.columns else comms
    df["Severity"] = df["Severity"] if "Severity" in df.columns else sevs
    df["Escalation status"] = df["Escalation status"] if "Escalation status" in df.columns else escs
    df["Relevant Team"] = df["Submitter LOB"]
    return df


def apply_optional_pending_review_updates(base_df, reviewed_df):
    """
    Applies manual updates from the optional Pending_Cases_Review sheet onto the Raw Data.
    If the user changes any ticket status (Reasonable <-> Unreasonable <-> SNULL) or edits
    Comment / Severity / Escalation status / Relevant Team, it updates the main dataset for final reporting.
    """
    df = base_df.copy()
    rev = reviewed_df.copy()

    res_col = "The responsible person replies with the handling result."
    loop_col = "Has the responsible person promptly closed the loop?"

    df["Order No."] = df["Order No."].astype(str).str.strip()
    rev["Order No."] = rev["Order No."].astype(str).str.strip()
    rev = rev.drop_duplicates(subset=["Order No."], keep="last")

    if "Escalation status " in rev.columns and "Escalation status" not in rev.columns:
        rev["Escalation status"] = rev["Escalation status "]

    cols_to_sync = [
        res_col,
        loop_col,
        "Comment",
        "Severity",
        "Escalation status",
        "Relevant Team",
        "Submitter LOB",
        "Submitter Name"
    ]
    for col in cols_to_sync:
        if col in rev.columns:
            valid_rev = rev[rev[col].notna() & (rev[col].astype(str).str.strip() != "") & (rev[col].astype(str).str.strip() != "nan")]
            update_map = dict(zip(valid_rev["Order No."], valid_rev[col].astype(str).str.strip()))
            if col in df.columns:
                df[col] = df["Order No."].map(update_map).fillna(df[col])
            else:
                df[col] = df["Order No."].map(update_map).fillna("")

    return df


def create_styled_excel(sheets_dict, add_dropdowns_on_sheet=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    editable_header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    regular_font = Font(name="Calibri", size=11, color="1E293B")
    bold_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    link_font = Font(name="Calibri", size=10, color="2563EB", underline="single")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    editable_cols = {
        "Comment", "Severity", "Escalation status", "Relevant Team",
        "The responsible person replies with the handling result."
    }

    for sheet_name, df in sheets_dict.items():
        ws = wb.create_sheet(title=sheet_name)
        cols = list(df.columns)

        for col_idx, col_name in enumerate(cols, start=1):
            cell = ws.cell(1, col_idx, col_name)
            cell.fill = editable_header_fill if (sheet_name == add_dropdowns_on_sheet and col_name in editable_cols) else header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for r_idx, row in enumerate(df.itertuples(index=False), start=2):
            for c_idx, val in enumerate(row, start=1):
                col_name = cols[c_idx - 1]
                clean_v = "" if pd.isna(val) else val
                cell = ws.cell(r_idx, c_idx, clean_v)

                # Make Electronic Flow Link a clickable hyperlink in Excel!
                if col_name == "Electronic Flow Link" and str(clean_v).startswith("http"):
                    cell.hyperlink = str(clean_v)
                    cell.font = link_font
                else:
                    cell.font = bold_font if c_idx == 1 else regular_font

                if r_idx % 2 == 1:
                    cell.fill = zebra_fill
                cell.border = thin_border
                if col_name in ["Comment", "Electronic Flow Link", "Issue Type", "Responsible Person's Reply and Comments", "E-Flow Problem Description"]:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=(col_name in ["Comment", "Responsible Person's Reply and Comments", "E-Flow Problem Description"]))
                else:
                    cell.alignment = Alignment(horizontal="center" if c_idx > 1 else "left", vertical="center")

        ws.freeze_panes = "A2"
        if len(df) > 0:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df) + 1}"

        for c_idx, col_name in enumerate(cols, start=1):
            col_letter = get_column_letter(c_idx)
            if col_name in ["Comment", "Responsible Person's Reply and Comments", "E-Flow Problem Description"]:
                ws.column_dimensions[col_letter].width = 42
            elif col_name in ["Issue Type", "Has the responsible person promptly closed the loop?"]:
                ws.column_dimensions[col_letter].width = 32
            elif col_name in ["Submitter Name", "Electronic Flow Link"]:
                ws.column_dimensions[col_letter].width = 28
            else:
                ws.column_dimensions[col_letter].width = 20

        if sheet_name == add_dropdowns_on_sheet and len(df) > 0:
            max_r = len(df) + 1
            res_col_name = "The responsible person replies with the handling result."
            if res_col_name in cols:
                r_col = get_column_letter(cols.index(res_col_name) + 1)
                dv_res = DataValidation(type="list", formula1='"Reasonable,Unreasonable,SNULL"', allow_blank=True)
                ws.add_data_validation(dv_res)
                dv_res.add(f"{r_col}2:{r_col}{max_r}")
            if "Severity" in cols:
                sev_col = get_column_letter(cols.index("Severity") + 1)
                dv_sev = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
                ws.add_data_validation(dv_sev)
                dv_sev.add(f"{sev_col}2:{sev_col}{max_r}")
            if "Escalation status" in cols:
                esc_col = get_column_letter(cols.index("Escalation status") + 1)
                dv_esc = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
                ws.add_data_validation(dv_esc)
                dv_esc.add(f"{esc_col}2:{esc_col}{max_r}")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def render_report(agents_df, hc_df):
    # --- Main UI Header (Exact User Code) ---
    st.title("⚡ Smart Feedback Analyzer & Fast Link-Review Workspace")

    col_u1, col_u2, col_u3 = st.columns(3)
    with col_u1:
        sys_file = st.file_uploader(
            "📂 1. Upload System Export Sheet (`明细表_Schedules_en.xlsx` or Raw Data):",
            type=["xlsx", "xls", "csv"],
            key="sys_uploader"
        )
    with col_u2:
        prev_file = st.file_uploader(
            "🧠 2. (Optional) Upload Previous Analyzed Sheet to Auto-Remember Old Orders:",
            type=["xlsx", "xls"],
            key="prev_uploader"
        )
    with col_u3:
        optional_pending_review_file = st.file_uploader(
            "📝 3. (Optional) Upload Updated 'Pending_Cases_Review.xlsx' After Manual Review:",
            type=["xlsx", "xls", "csv"],
            key="pending_review_uploader"
        )

    if sys_file is not None:
        # Build historical memory map if previous analyzed file is provided
        hist_map = {}
        if prev_file is not None:
            xls_prev = pd.ExcelFile(prev_file)
            for sname in xls_prev.sheet_names:
                p_df = pd.read_excel(xls_prev, sheet_name=sname)
                res_c = "The responsible person replies with the handling result."
                if "Order No." in p_df.columns and res_c in p_df.columns:
                    for _, prow in p_df.iterrows():
                        ono = str(prow["Order No."]).strip()
                        rval = str(prow[res_c]).strip()
                        if ono and rval in ["Reasonable", "Unreasonable", "SNULL"]:
                            hist_map[ono] = {
                                "result": rval,
                                "comment": str(prow.get("Comment", "")) if pd.notna(prow.get("Comment", "")) else "",
                                "severity": str(prow.get("Severity", "")) if pd.notna(prow.get("Severity", "")) else "",
                                "escalation": str(prow.get("Escalation status", "")) if pd.notna(prow.get("Escalation status", "")) else ""
                            }
            st.info(f"🧠 Loaded `{len(hist_map)}` previously reviewed Order Nos from historical sheet!")

        # Read system file (without automatic multi-tab background sync unless explicitly uploaded)
        if sys_file.name.endswith(".csv"):
            raw_sys_df = pd.read_csv(sys_file)
        else:
            xls_sys = pd.ExcelFile(sys_file)
            target_s = (
                "Schedule of Details"
                if "Schedule of Details" in xls_sys.sheet_names
                else ("Raw Data" if "Raw Data" in xls_sys.sheet_names else ("Enriched_Raw_Data" if "Enriched_Raw_Data" in xls_sys.sheet_names else xls_sys.sheet_names[0]))
            )
            raw_sys_df = pd.read_excel(xls_sys, sheet_name=target_s)
            if "Consignee" not in [str(c).strip() for c in raw_sys_df.columns]:
                for s_cand in xls_sys.sheet_names:
                    df_cand = pd.read_excel(xls_sys, sheet_name=s_cand)
                    if "Consignee" in [str(c).strip() for c in df_cand.columns]:
                        raw_sys_df = df_cand
                        break

        enriched_df = enrich_and_classify_dataframe(raw_sys_df, agents_df, hist_map)

        # Apply updates from Optional Updated Pending_Cases_Review file if uploaded
        if optional_pending_review_file is not None:
            if optional_pending_review_file.name.endswith(".csv"):
                rev_up_df = pd.read_csv(optional_pending_review_file)
            else:
                xls_rev_up = pd.ExcelFile(optional_pending_review_file)
                rev_s_name = "Pending_Cases_Review" if "Pending_Cases_Review" in xls_rev_up.sheet_names else xls_rev_up.sheet_names[0]
                rev_up_df = pd.read_excel(xls_rev_up, sheet_name=rev_s_name)

            if "Order No." in rev_up_df.columns:
                enriched_df = apply_optional_pending_review_updates(enriched_df, rev_up_df)
                st.success(f"✅ Applied manual updates from `{optional_pending_review_file.name}`! Any status changes (Reasonable ↔ Unreasonable / SNULL) are now calculated in the final report.")
            else:
                st.error("❌ The uploaded Pending_Cases_Review sheet must contain an 'Order No.' column.")

        # Check for unmapped Consignee IDs and let user map them right here!
        unmapped_mask = (
            enriched_df["Submitter Name"].isna()
            | (enriched_df["Submitter Name"] == "Unknown Agent")
            | (enriched_df["Submitter Name"] == "nan")
        )
        if unmapped_mask.any():
            unmapped_ids = enriched_df.loc[unmapped_mask, "Consignee"].unique().tolist()
            with st.expander(f"⚠️ {len(unmapped_ids)} New Consignee IDs Found! Click here to assign their Submitter Name & LOB once:", expanded=True):
                new_roster_rows = pd.DataFrame({
                    "Consignee": unmapped_ids,
                    "Submitter Name": ["" for _ in unmapped_ids],
                    "Submitter LOB": [hc_df["LOB"].iloc[0] if len(hc_df) > 0 else "Egypt Hotline" for _ in unmapped_ids]
                })
                edited_new_agents = st.data_editor(
                    new_roster_rows,
                    column_config={
                        "Submitter LOB": st.column_config.SelectboxColumn("Submitter LOB", options=sorted(hc_df["LOB"].unique().tolist()))
                    },
                    use_container_width=True,
                    key="quick_map_editor"
                )
                if st.button("✅ Save New Agents to Master Roster & Re-Map"):
                    valid_new = edited_new_agents[edited_new_agents["Submitter Name"].astype(str).str.strip() != ""]
                    if len(valid_new) > 0:
                        combined = pd.concat([agents_df, valid_new], ignore_index=True).drop_duplicates(subset=["Consignee"], keep="last")
                        save_master_roster(combined, hc_df)
                        st.rerun()

        # Show Smart Pre-Classification Summary Banner
        res_col = "The responsible person replies with the handling result."
        auto_counts = enriched_df[res_col].value_counts().to_dict()
        b1, b2, b3, b4, b5 = st.columns(5)
        b1.metric("Total Uploaded Orders", len(enriched_df))
        b2.metric("✅ Auto-Detected Reasonable", auto_counts.get("Reasonable", 0))
        b3.metric("❌ Auto-Detected Unreasonable", auto_counts.get("Unreasonable", 0))
        b4.metric("⏳ Auto-Detected SNULL (Pending)", auto_counts.get("SNULL", 0))
        b5.metric("🔍 Needs Quick Link Check", auto_counts.get("Needs Link Check", 0))

        st.divider()

        # --- Interactive Fast Review Workspace (Stage 1 & Stage 2 Combined or Separate Excel) ---
        st.subheader("🔍 Step 2: Fast Review Workspace (Click Link, Verify & Edit In-Place or Export to Excel)")
        filter_choice = st.radio(
            "Filter Cases to Review:",
            options=[
                "🔍 Needs Link Check + Unreasonable + SNULL (Recommended - Skip Auto-Reasonable)",
                "🔎 Only Ambiguous Cases ('Needs Link Check')",
                "📋 All Uploaded Cases"
            ],
            horizontal=True
        )

        if "Only Ambiguous" in filter_choice:
            view_mask = enriched_df[res_col] == "Needs Link Check"
        elif "Recommended" in filter_choice:
            view_mask = enriched_df[res_col].isin(["Needs Link Check", "Unreasonable", "SNULL"])
        else:
            view_mask = pd.Series([True] * len(enriched_df), index=enriched_df.index)

        review_cols = [
            "Order No.", "Bill of Lading Time", "Consignee", "Submitter Name", "Submitter LOB",
            "Issue Type", "Knowledge ID", "Responsible person replies with a handling solution",
            "Is it resolved?", "Has the responsible person promptly closed the loop?",
            "Responsible Person's Reply and Comments", "Electronic Flow Link",
            "Auto_Detection_Confidence", res_col, "Comment", "Severity", "Escalation status", "Relevant Team"
        ]
        review_cols = [c for c in review_cols if c in enriched_df.columns]

        edited_review_df = st.data_editor(
            enriched_df.loc[view_mask, review_cols],
            column_config={
                "Electronic Flow Link": st.column_config.LinkColumn(
                    "Electronic Flow Link",
                    help="Click to open the case directly in Huawei Knowledge System",
                    display_text="🔗 Open E-Flow Link"
                ),
                res_col: st.column_config.SelectboxColumn(
                    "Handling Result (Select)",
                    options=["Reasonable", "Unreasonable", "SNULL", "Needs Link Check"],
                    required=True
                ),
                "Severity": st.column_config.SelectboxColumn(
                    "Severity",
                    options=["", "High", "Medium", "Low"]
                ),
                "Escalation status": st.column_config.SelectboxColumn(
                    "Escalation status",
                    options=["Yes", "No"]
                ),
                "Relevant Team": st.column_config.SelectboxColumn(
                    "Relevant Team (=LOB)",
                    options=sorted(list(set(hc_df["LOB"].tolist() + enriched_df["Submitter LOB"].unique().tolist())))
                )
            },
            use_container_width=True,
            height=420,
            key="fast_workspace_editor"
        )

        # Sync edits from the interactive table back into enriched_df
        enriched_df.loc[edited_review_df.index, review_cols] = edited_review_df

        # Treat any remaining 'Needs Link Check' as 'SNULL' for standard 3-category pivot math, or let user see
        final_working_df = enriched_df.copy()
        final_working_df[res_col] = final_working_df[res_col].replace("Needs Link Check", "SNULL")

        # Build Pending_Cases_Review (19-column exact layout for Unreasonable & SNULL)
        pending_mask = final_working_df[res_col].isin(["Unreasonable", "SNULL"])
        p_sub = final_working_df[pending_mask].copy()
        sol_col_1 = "Responsible person replies with a handling solution"
        sol_col_2 = "Responsible person replies with a handling solution.1" if "Responsible person replies with a handling solution.1" in p_sub.columns else sol_col_1

        pending_review_19_df = pd.DataFrame({
            "Order No.": p_sub.get("Order No.", ""),
            "Bill of Lading Time": p_sub.get("Bill of Lading Time", ""),
            "Issue Type": p_sub.get("Issue Type", ""),
            "Country/Region Where the Issue Occurred": p_sub.get("Country/Region Where the Issue Occurred", ""),
            "Process Type": p_sub.get("Process Type", ""),
            "Consignee": p_sub.get("Consignee", ""),
            "Submitter Name": p_sub.get("Submitter Name", ""),
            "Submitter LOB": p_sub.get("Submitter LOB", ""),
            "Knowledge ID": p_sub.get("Knowledge ID", ""),
            "Responsible person replies with a handling solution (Initial)": p_sub.get(sol_col_1, "SNULL"),
            "Has the responsible person promptly closed the loop?": p_sub.get("Has the responsible person promptly closed the loop?", ""),
            "Person Responsible": p_sub.get("Person Responsible", "SNULL"),
            "Responsible person replies with a handling solution (Final)": p_sub.get(sol_col_2, "SNULL"),
            "The responsible person replies with the handling result.": p_sub.get(res_col, "SNULL"),
            "Electronic Flow Link": p_sub.get("Electronic Flow Link", ""),
            "Comment": p_sub.get("Comment", ""),
            "Severity": p_sub.get("Severity", ""),
            "Escalation status": p_sub.get("Escalation status", "Yes"),
            "Relevant Team": p_sub.get("Relevant Team", p_sub.get("Submitter LOB", ""))
        })

        # Stage 1 Download Button (if user prefers editing in Excel first and re-uploading)
        stage1_bytes = create_styled_excel(
            {
                "Pending_Cases_Review": pending_review_19_df,
                "Enriched_Raw_Data": final_working_df
            },
            add_dropdowns_on_sheet="Pending_Cases_Review"
        )
        st.download_button(
            label="📥 Stage 1 Option: Download 'Pending_Cases_Review.xlsx' (With Clickable Links & Auto-Pre-Filled Results) for Offline Excel Review",
            data=stage1_bytes,
            file_name="Pending_Cases_Review.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )

        st.divider()

        # ==========================================================================
        # LIVE FINAL ANALYSIS, CHARTS & PDF / EXCEL REPORT (BOTH OLD & NEW CLOSURE)
        # ==========================================================================
        st.subheader("📊 Step 3: Live Final Analysis & Executive PDF / Excel Report")
        loop_col = "Has the responsible person promptly closed the loop?"
        all_lobs = sorted(list(set(hc_df["LOB"].tolist() + final_working_df["Submitter LOB"].unique().tolist())))
        hc_map = dict(zip(hc_df["LOB"], hc_df["Team HC"]))

        summary_rows, pivot_rows = [], []
        for lob in all_lobs:
            sub = final_working_df[final_working_df["Submitter LOB"] == lob]
            hc = int(hc_map.get(lob, max(1, len(sub["Submitter Name"].unique()))))
            part = sub["Submitter Name"].nunique() if len(sub) > 0 else 0
            total_tix = len(sub)
            usage_rate = (part / hc * 100) if hc > 0 else 0.0

            reas_cnt = int((sub[res_col] == "Reasonable").sum())
            snull_cnt = int((sub[res_col] == "SNULL").sum())
            unreas_cnt = int((sub[res_col] == "Unreasonable").sum())

            # Old Closure Status (Loop-Based): Timely closure of the loop + Not through the owner phase
            closed_old_cnt = int(sub[loop_col].isin(["Timely closure of the loop", "Not through the owner phase"]).sum())
            clos_old_rate = (closed_old_cnt / total_tix * 100) if total_tix > 0 else None

            # New Closure Status (Handling-Result-Based): (Reasonable + Unreasonable) / Total, and 100% if Total == 0
            closed_new_cnt = reas_cnt + unreas_cnt
            clos_new_rate = (closed_new_cnt / total_tix * 100) if total_tix > 0 else 100.0

            reas_rate = (reas_cnt / total_tix * 100) if total_tix > 0 else None

            summary_rows.append({
                "LOB": lob,
                "Team HC": hc,
                "Participation": part,
                "Usage Rate per Team (%)": round(usage_rate, 1),
                "Reasonability (%)": round(reas_rate, 1) if reas_rate is not None else None,
                "Closure Status (Old - Loop %)": round(clos_old_rate, 1) if clos_old_rate is not None else None,
                "Closure Status (New - Result %)": round(clos_new_rate, 1),
                "Total Feedbacks": total_tix
            })

            if total_tix > 0:
                pivot_rows.append({
                    "Row Labels": lob,
                    "Reasonable": reas_cnt,
                    "SNULL": snull_cnt,
                    "Unreasonable": unreas_cnt,
                    "Grand Total": total_tix,
                    "Reasonability (%)": round(reas_rate, 1),
                    "Closure Status (Old - Loop %)": round(clos_old_rate, 1) if clos_old_rate is not None else 0.0,
                    "Closure Status (New - Result %)": round(clos_new_rate, 1)
                })

        summary_df_out = pd.DataFrame(summary_rows)
        pivot_df_out = pd.DataFrame(pivot_rows)

        total_hc_all = int(summary_df_out["Team HC"].sum())
        total_part_all = int(final_working_df["Submitter Name"].nunique())
        total_tix_all = len(final_working_df)
        total_reas_all = int((final_working_df[res_col] == "Reasonable").sum())
        total_closed_old_all = int(final_working_df[loop_col].isin(["Timely closure of the loop", "Not through the owner phase"]).sum())
        total_closed_new_all = int((final_working_df[res_col] != "SNULL").sum())
        escalated_cnt = int((pending_review_19_df["Escalation status"].astype(str).str.strip().str.upper() == "YES").sum())

        k1, k2, k3, k4, k5, k6, k7 = st.columns(7)
        k1.metric("Total Feedbacks", total_tix_all)
        k2.metric("Active Participants", f"{total_part_all} / {total_hc_all}")
        k3.metric("Overall Usage Rate", f"{(total_part_all/total_hc_all*100):.1f}%" if total_hc_all else "0%")
        k4.metric("Reasonability Rate", f"{(total_reas_all/total_tix_all*100):.1f}%" if total_tix_all else "0%")
        k5.metric("Closure (Old - Loop)", f"{(total_closed_old_all/total_tix_all*100):.1f}%" if total_tix_all else "0%")
        k6.metric("Closure (New - Result)", f"{(total_closed_new_all/total_tix_all*100):.1f}%" if total_tix_all else "0%")
        k7.metric("Pending / Unreasonable", len(pending_review_19_df))

        ch1, ch2 = st.columns(2)
        with ch1:
            active_sum = summary_df_out[summary_df_out["Total Feedbacks"] > 0]
            fig_bar = px.bar(
                active_sum,
                x="LOB",
                y=[
                    "Usage Rate per Team (%)",
                    "Reasonability (%)",
                    "Closure Status (Old - Loop %)",
                    "Closure Status (New - Result %)"
                ],
                barmode="group",
                title="Team Performance Comparison (Usage vs. Reasonability vs. Old & New Closure %)",
                color_discrete_sequence=["#2563eb", "#10b981", "#14b8a6", "#8b5cf6"]
            )
            st.plotly_chart(fig_bar, use_container_width=True)
        with ch2:
            fig_pie = px.pie(
                final_working_df,
                names=res_col,
                title="Overall Handling Result Distribution",
                color=res_col,
                color_discrete_map={"Reasonable": "#10b981", "Unreasonable": "#ef4444", "SNULL": "#f59e0b"}
            )
            st.plotly_chart(fig_pie, use_container_width=True)

        t_sum_tab, t_pend_tab, t_piv_tab = st.tabs([
            "📊 Team Summary Table",
            "⏳ Pending_Cases_Review (Unreasonable & SNULL)",
            "📈 Handling Result Pivot"
        ])
        with t_sum_tab:
            st.dataframe(summary_df_out, use_container_width=True)
        with t_pend_tab:
            st.dataframe(pending_review_19_df, use_container_width=True)
        with t_piv_tab:
            st.dataframe(pivot_df_out, use_container_width=True)

        # PDF Generator Function
        def generate_final_pdf_bytes():
            buf = io.BytesIO()
            doc = SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
            styles = getSampleStyleSheet()
            story = []

            title_style = ParagraphStyle("ExecTitle", parent=styles["Heading1"], fontSize=18, textColor=colors.HexColor("#0F172A"), spaceAfter=6)
            sub_style = ParagraphStyle("ExecSub", parent=styles["Normal"], fontSize=9.5, textColor=colors.HexColor("#334155"), spaceAfter=12)
            cell_wrap_style = ParagraphStyle("CellWrap", parent=styles["Normal"], fontSize=7, leading=8.5, textColor=colors.HexColor("#1E293B"))

            story.append(Paragraph("Executive Feedback Analysis Report", title_style))
            story.append(Paragraph(
                f"<b>Total Feedbacks:</b> {total_tix_all} &nbsp;|&nbsp; "
                f"<b>Active Participants:</b> {total_part_all} / {total_hc_all} ({(total_part_all/total_hc_all*100 if total_hc_all else 0):.1f}% Usage Rate) &nbsp;|&nbsp; "
                f"<b>Overall Reasonability:</b> {(total_reas_all/total_tix_all*100 if total_tix_all else 0):.1f}% &nbsp;|&nbsp; "
                f"<b>Closure Rate (Old Loop):</b> {(total_closed_old_all/total_tix_all*100 if total_tix_all else 0):.1f}% &nbsp;|&nbsp; "
                f"<b>Closure Rate (New Result):</b> {(total_closed_new_all/total_tix_all*100 if total_tix_all else 0):.1f}% &nbsp;|&nbsp; "
                f"<b>Pending/Unreasonable Cases:</b> {len(pending_review_19_df)} (Escalated: {escalated_cnt})",
                sub_style
            ))

            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(11.2, 3.5), gridspec_kw={"width_ratios": [2.2, 1]})
            active_s = summary_df_out[summary_df_out["Total Feedbacks"] > 0].fillna(0)
            if len(active_s) == 0:
                active_s = summary_df_out.fillna(0)
            x = range(len(active_s))
            w = 0.20
            ax1.bar([i - 1.5*w for i in x], active_s["Usage Rate per Team (%)"], width=w, label="Usage Rate (%)", color="#2563eb")
            ax1.bar([i - 0.5*w for i in x], active_s["Reasonability (%)"], width=w, label="Reasonability (%)", color="#10b981")
            ax1.bar([i + 0.5*w for i in x], active_s["Closure Status (Old - Loop %)"], width=w, label="Closure Old (Loop %)", color="#14b8a6")
            ax1.bar([i + 1.5*w for i in x], active_s["Closure Status (New - Result %)"], width=w, label="Closure New (Result %)", color="#8b5cf6")
            ax1.set_xticks(list(x))
            ax1.set_xticklabels(active_s["LOB"], rotation=30, ha="right", fontsize=8)
            ax1.set_ylim(0, 115)
            ax1.set_ylabel("Percentage (%)")
            ax1.set_title("Active Teams Performance Breakdown (%)", fontsize=10, fontweight="bold")
            ax1.legend(fontsize=7.5, loc="upper right", ncol=2)
            ax1.grid(axis="y", linestyle="--", alpha=0.4)

            res_counts = final_working_df[res_col].value_counts()
            pie_colors = [{"Reasonable": "#10b981", "Unreasonable": "#ef4444", "SNULL": "#f59e0b"}.get(k, "#64748b") for k in res_counts.index]
            ax2.pie(res_counts.values, labels=res_counts.index, autopct="%1.1f%%", startangle=140, colors=pie_colors, textprops={"fontsize": 8.5})
            ax2.set_title("Handling Result Distribution", fontsize=10, fontweight="bold")

            plt.tight_layout()
            img_buf = io.BytesIO()
            plt.savefig(img_buf, format="png", dpi=160)
            plt.close(fig)
            img_buf.seek(0)
            story.append(Image(img_buf, width=760, height=225))
            story.append(Spacer(1, 10))

            story.append(Paragraph("1. Team Summary & KPIs Table", styles["Heading2"]))
            sum_table_data = [[
                "LOB (Team)", "Team HC", "Participation", "Usage Rate (%)",
                "Reasonability (%)", "Closure Old (Loop %)", "Closure New (Result %)", "Total Feedbacks"
            ]]
            for _, r in summary_df_out.iterrows():
                sum_table_data.append([
                    str(r["LOB"]), str(r["Team HC"]), str(r["Participation"]),
                    f"{r['Usage Rate per Team (%)']:.1f}%",
                    f"{r['Reasonability (%)']:.1f}%" if pd.notna(r["Reasonability (%)"]) else "-",
                    f"{r['Closure Status (Old - Loop %)']:.1f}%" if pd.notna(r["Closure Status (Old - Loop %)"]) else "-",
                    f"{r['Closure Status (New - Result %)']:.1f}%" if pd.notna(r["Closure Status (New - Result %)"]) else "-",
                    str(r["Total Feedbacks"])
                ])
            sum_table_data.append([
                "Grand Total", str(total_hc_all), str(total_part_all),
                f"{(total_part_all/total_hc_all*100 if total_hc_all else 0):.1f}%",
                f"{(total_reas_all/total_tix_all*100 if total_tix_all else 0):.1f}%",
                f"{(total_closed_old_all/total_tix_all*100 if total_tix_all else 0):.1f}%",
                f"{(total_closed_new_all/total_tix_all*100 if total_tix_all else 0):.1f}%",
                str(total_tix_all)
            ])

            t_sum = Table(sum_table_data, colWidths=[135, 58, 72, 85, 95, 105, 110, 85], repeatRows=1)
            t_sum.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            story.append(t_sum)
            story.append(PageBreak())

            story.append(Paragraph("2. Pending & Unreasonable Cases (Pending_Cases_Review)", styles["Heading2"]))
            act_data = [["Order No.", "Submitter Name", "Issue Type", "Knowledge ID", "Result", "Severity", "Escalated", "Relevant Team", "Review Comment / Action Taken"]]
            for _, r in pending_review_19_df.iterrows():
                comment_txt = str(r.get("Comment", "") if pd.notna(r.get("Comment", "")) else "").strip()
                if comment_txt == "nan":
                    comment_txt = ""
                act_data.append([
                    str(r.get("Order No.", "")),
                    Paragraph(str(r.get("Submitter Name", "")), cell_wrap_style),
                    Paragraph(str(r.get("Issue Type", "")), cell_wrap_style),
                    str(r.get("Knowledge ID", "")),
                    str(r.get(res_col, "")),
                    str(r.get("Severity", "")) if pd.notna(r.get("Severity", "")) else "-",
                    str(r.get("Escalation status", "")) if pd.notna(r.get("Escalation status", "")) else "-",
                    Paragraph(str(r.get("Relevant Team", "")), cell_wrap_style),
                    Paragraph(comment_txt[:260], cell_wrap_style)
                ])

            t_act = Table(act_data, colWidths=[82, 95, 105, 68, 62, 48, 48, 85, 195], repeatRows=1)
            t_act.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (3, 0), (6, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ]))
            story.append(t_act)

            doc.build(story)
            buf.seek(0)
            return buf.getvalue()

        dl1, dl2 = st.columns(2)
        with dl1:
            st.download_button(
                label="📄 Download Final Executive PDF Report (100% English + Charts + KPIs)",
                data=generate_final_pdf_bytes(),
                file_name="Final_Feedback_Analysis_Report.pdf",
                mime="application/pdf",
                type="primary",
                use_container_width=True
            )
        with dl2:
            final_wb_bytes = create_styled_excel(
                {
                    "Summary": summary_df_out,
                    "Handling_Result_Pivot": pivot_df_out,
                    "Pending_Cases_Review": pending_review_19_df,
                    "Enriched_Raw_Data": final_working_df
                },
                add_dropdowns_on_sheet="Pending_Cases_Review"
            )
            st.download_button(
                label="📥 Download Final Complete Excel Workbook (.xlsx)",
                data=final_wb_bytes,
                file_name="Final_Feedback_Analysis_Workbook.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
